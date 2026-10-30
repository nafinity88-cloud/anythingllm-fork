from pathlib import Path
import re, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sources import SOURCES

ROOT=Path(__file__).resolve().parents[1]
source_map={x[0]:i+1 for i,x in enumerate(SOURCES)}
raw=(ROOT/'literature'/'report.md').read_text(encoding='utf-8')
def refs(s):
    return re.sub(r'\{([A-Z0-9_]+)\}',lambda m:'['+str(source_map[m[1]])+']',s)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.68)
sec.left_margin=sec.right_margin=Inches(.75)
sec.header_distance=sec.footer_distance=Inches(.3)
for st in doc.styles:
    for e in list(st.element.iter(qn('w:pBdr'))): e.getparent().remove(e)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.1
normal.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(12)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(6)
header=sec.header.paragraphs[0];header.text='BINTANG MONO INDONESIA  |  RISET KOMPETITOR'
header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Sumber publik  •  30 Oktober 2026  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)

def rich(para,s):
    for i,x in enumerate(re.split(r'\*\*(.*?)\*\*',refs(s))):
        run=para.add_run(x);run.bold=bool(i%2)
def para(s):
    p=doc.add_paragraph();rich(p,s);return p
def table(lines,page_no):
    rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    rows=[row for row in rows if not all(re.fullmatch(r'[-: ]+',x) for x in row)]
    count=len(rows[0]); t=doc.add_table(rows=0,cols=count);t.autofit=False
    if count==4:
        widths=[1.1,1.6,1.95,2.35] if page_no==3 else [1.15,1.8,1.95,2.1]
    elif page_no==1:widths=[1.65,1.85,3.5]
    elif page_no==5:widths=[1.45,2.8,2.75]
    elif page_no==6:widths=[2.7,2.15,2.15]
    else:widths=[1.45,2.9,2.65]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        pr=cells[0]._tc.getparent().get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for j,txt in enumerate(row):
            cells[j].width=Inches(widths[j]);p=cells[j].paragraphs[0];rich(p,txt)
            p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(4)
            p.paragraph_format.line_spacing=1.05
            for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEFF2');cells[j]._tc.get_or_add_tcPr().append(sh)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0)

for pi,page in enumerate(raw.split('---PAGE---')):
    if pi:doc.add_page_break()
    lines=page.strip().splitlines();i=0
    while i<len(lines):
        ln=lines[i].strip()
        if not ln:i+=1;continue
        if ln.startswith('|'):
            tb=[]
            while i<len(lines) and lines[i].strip().startswith('|'):tb.append(lines[i]);i+=1
            table(tb,pi);continue
        if ln.startswith('# '):
            if pi==0:doc.add_paragraph(ln[2:],'Title')
            else:doc.add_heading(ln[2:],1)
        elif ln.startswith('## '):doc.add_heading(ln[3:],2)
        else:para(ln)
        i+=1

def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz)
    color=OxmlElement('w:color');color.set(qn('w:val'),'245575');pr.append(color);r.append(pr)
    tx=OxmlElement('w:t');tx.text=label;r.append(tx);el.append(r);p._p.append(el)

for start in range(0,len(SOURCES),13):
    doc.add_page_break();doc.add_heading('Daftar sumber' + (' lanjutan' if start else ''),1)
    para('Nomor pada isi mengacu pada daftar ini. Semua halaman diakses atau ditelusuri pada 30 Oktober 2026. Keterangan perusahaan dan spesifikasi produsen belum menggantikan verifikasi izin aktif atau pengujian independen.')
    for i,(key,owner,title,url,note) in enumerate(SOURCES[start:start+13],start+1):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
        p.add_run(f'[{i}] {owner} — ').bold=True
        hyperlink(p,title,url);p.add_run('. '+note)
        for r in p.runs:r.font.size=Pt(10)
        p.paragraph_format.keep_together=True

doc.core_properties.title='Riset kompetitor Bintang Mono Indonesia'
doc.core_properties.subject='Rapid test analyzer dan tabung darah di Indonesia'
doc.core_properties.author='Codex'
output=ROOT/'to_human'/'Riset Mendalam Kompetitor BMI.docx';doc.save(output)
(ROOT/'data'/'sources.json').write_text(json.dumps([dict(id=i+1,key=k,publisher=o,title=t,url=u,note=n) for i,(k,o,t,u,n) in enumerate(SOURCES)],ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'findings.md').write_text(refs(raw),encoding='utf-8')
print('Sources:',len(SOURCES));print('Words:',len(raw.split()));print('Logical pages:',len(raw.split('---PAGE---'))+3);print(output)
