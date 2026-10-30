from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'to_human' / 'Riset Bintang Mono Indonesia.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(11)
    doc.styles[name].paragraph_format.space_after = Pt(7)
header = sec.header.paragraphs[0]
header.text = 'BINTANG MONO INDONESIA  |  LAPORAN RISET PERUSAHAAN'
header.runs[0].font.size = Pt(8)
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Riset sumber publik  •  30 Oktober 2026  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
for r in foot.runs: r.font.size = Pt(8)

alltext=[]
def p(t,style=None):
    alltext.append(t)
    return doc.add_paragraph(t,style)
def h(t,level=2):
    alltext.append(t); doc.add_heading(t,level)
def page(title):
    doc.add_page_break(); h(title,1)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    if widths:
        for col, width in zip(t.columns, widths): col.width=Inches(width)
    for i,s in enumerate(headers): t.rows[0].cells[i].text=s
    for row in rows:
        cells=t.add_row().cells
        for i,s in enumerate(row): cells[i].text=s
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if ri==0: pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[ci])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(5)
                para.paragraph_format.space_before=Pt(4)
                for r in para.runs:
                    r.font.size=Pt(10)
                    if ri==0: r.bold=True
            if ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EDEFF2'); c._tc.get_or_add_tcPr().append(sh)
    p('')

p('Riset Bintang Mono Indonesia','Title')
p('Profil usaha dan peluang pengembangan berbasis sumber publik','Subtitle')
p('Laporan riset perusahaan  |  30 Oktober 2026  |  Bahasa Indonesia')
h('Ringkasan eksekutif',1)
p('PT Bintang Mono Indonesia atau BMI memiliki jejak publik yang konsisten sebagai pemasok alat kesehatan, khususnya diagnostik in vitro dan kardiovaskular. Situs perusahaan dan profil LinkedIn menampilkan bidang usaha serta alamat yang selaras. Publikasi resmi DJKI juga mencantumkan perusahaan sebagai pemohon merek Monotes dengan alamat yang sama. Bukti ini mendukung identitas bisnis dan keterkaitannya dengan merek, tetapi belum mengukur skala atau kinerja usahanya. [1][2][3][10]')
p('Katalog daring menunjukkan cakupan tabung pengambilan darah, rapid test penyakit infeksi, analyzer, serta peralatan dan bahan habis pakai yang ditempatkan pada kategori produk dalam negeri. Kombinasi tersebut membuka hipotesis penjualan lintas kategori kepada pelanggan laboratorium dan fasilitas kesehatan. Peluang pembelian berulang perlu diuji dengan data transaksi, pemakaian, dan stok aktual. [4][5][6][7]')
p('Prioritas yang disarankan adalah merapikan data produk beserta dokumen pendukung, menguji paket penawaran per segmen, dan mengukur pembelian ulang. Penilaian kelayakan kerja sama tetap memerlukan konfirmasi izin aktif, kewenangan distribusi, kemampuan layanan, dan data keuangan. Informasi itu belum terbukti lengkap dari sumber yang diperiksa.')
h('Cara membaca laporan')
p('Fakta publik berarti informasi yang ditemukan pada sumber yang disebutkan. Keterangan perusahaan adalah pernyataan dari kanal milik perusahaan dan belum otomatis terverifikasi secara independen. Analisis dan rekomendasi merupakan penalaran peneliti yang masih perlu diuji.')
p('Cakupan laporan meliputi identitas, portofolio, kanal bisnis, bukti pendukung, kesenjangan informasi, serta usulan tindak lanjut. Riset ini merupakan penelusuran perusahaan berbasis sumber publik; tidak mencakup wawancara, kunjungan, audit, atau eksperimen klinis.')

page('Profil perusahaan dan arah bisnis')
h('Identitas dan kontak')
p('Nama yang digunakan pada kanal resmi adalah PT Bintang MONO Indonesia. Kantornya dicantumkan di Taman Tekno BSD Sektor XI Blok J 03 No 25, Setu, Tangerang Selatan, Banten 15314. Nomor kontak yang dipublikasikan adalah (021) 29040042, (021) 75684100, dan +62 811 795 577; alamat surel layanan pelanggan adalah customercare@ptbintangmono.com. [1]')
p('Profil LinkedIn perusahaan menyebut tahun pendirian 2008 dan ukuran perusahaan 11–50 karyawan. Kedua informasi ini bersifat pelaporan mandiri. Angka ukuran tersebut tidak diperlakukan sebagai jumlah pegawai aktual karena daftar pengguna yang mengaitkan dirinya dengan perusahaan dapat berbeda dan tidak selalu tervalidasi. [3]')
h('Fokus usaha dan merek')
p('BMI menyatakan berperan sebagai importir dan distributor alat kesehatan diagnostik in vitro serta kardiovaskular. Profil LinkedIn mencantumkan MONOTES, ORIENT GENE, dan HEALGEN untuk diagnostik; serta XLIMUS, RESTORE DEB, PRECISION, dan LEGLOW RX/OTW untuk kardiovaskular. Pencantuman itu menunjukkan portofolio yang dikomunikasikan perusahaan, bukan bukti masa berlaku penunjukan atau eksklusivitas saat ini. [2][3]')
h('Nilai dan posisi yang ingin dibangun')
p('Halaman Tentang Kami menempatkan hubungan dengan mitra, inovasi, fokus pelanggan, dan kemanusiaan sebagai nilai perusahaan. Misinya juga menyebut etika usaha, kepatuhan, dan nilai syariah universal. Ini berguna untuk memahami arah komunikasi perusahaan; pelaksanaan nilai tersebut memerlukan bukti operasional. [2]')
h('Implikasi bagi model bisnis')
p('Analisis: penyebutan subdistributor dan konsumen pada profil perusahaan menunjukkan pendekatan penjualan yang melibatkan mitra serta pembeli akhir. Kombinasi alat dan bahan habis pakai berpotensi menghasilkan hubungan pelanggan berulang, tetapi struktur pendapatan, jumlah mitra aktif, wilayah layanan, dan kontribusi tiap merek belum dapat dihitung. [2][4][6]')
p('Untuk menilai kekuatan model ini, data yang paling berguna adalah penjualan per kategori, frekuensi pemesanan ulang, pelanggan aktif per wilayah, margin setelah biaya distribusi, dan konsentrasi pembelian pada pelanggan terbesar.')

page('Peta portofolio dan kanal penjualan')
p('Daftar berikut merangkum contoh yang terlihat pada katalog publik. Kehadiran pada situs tidak memastikan stok, harga, spesifikasi lengkap, atau status izin setiap produk pada tanggal laporan.')
table(['Kelompok','Contoh pada sumber publik','Sumber'],[
('Pengambilan darah dan sampel','Vacuum tube EDTA, plain, gel dan clot activator, sodium citrate, micro EDTA, sample cup dan sample tube','[4]'),
('Rapid test penyakit infeksi','ORIENT GENE untuk HIV, HCV, HBsAg, malaria, dengue dan sifilis','[7]'),
('Analyzer imunofluoresensi','MONO X dan Mono Hba1c','[6]'),
('Kategori produk dalam negeri','Alcohol swab, centrifuge 8 dan 12 hole, laminar air flow, bio safety cabinet','[5]'),
('Pemantauan koagulasi','Hemonart CompactOne 9 disebut pada artikel perusahaan','[9]')
],[1.55,4.5,.8])
p('Label produk dalam negeri mengikuti pengelompokan situs. Laporan ini belum memeriksa pabrik aktual, skema produksi, nomor izin AKD masing-masing produk, atau dokumen tingkat komponen dalam negeri. Identitas produsen dan peran BMI perlu dibedakan per produk. [5]')
h('Kanal yang dapat diamati')
p('Halaman produk alcohol swab dan MONO X menyediakan tautan e-katalog. Artikel perusahaan mengenai Hemonart juga mengarahkan pembaca ke katalog Inaproc. Halaman Inaproc tidak berhasil dibaca melalui alat riset, sehingga harga, status tayang, dan ketersediaan terkini belum dikonfirmasi. Tautan Monoklik pada situs mengarah ke halaman login; proses pemesanan di dalamnya tidak dievaluasi. [8][9][12]')
p('PT Axion Jaya Alkesindo mencantumkan PT Bintang Mono Indonesia pada bagian mitra kerja sama di situsnya. Ini menjadi bukti dari pihak lain mengenai hubungan bisnis yang dikomunikasikan, tetapi tidak menjelaskan periode, nilai transaksi, atau ruang lingkup kerja sama. [11]')
h('Peluang yang perlu diuji')
p('Analisis: penawaran dapat dikelompokkan berdasarkan kebutuhan pelanggan, misalnya perlengkapan pengambilan sampel, pemeriksaan infeksi, atau alat beserta bahan pendukung. Sebelum membuat paket, pastikan kompatibilitas, izin, ketersediaan, dan kebutuhan pembeli; keberadaan kategori yang berdekatan belum membuktikan paket tertentu layak dijual.')

page('Kekuatan bukti dan kesenjangan informasi')
h('Temuan independen yang paling jelas')
p('Publikasi resmi DJKI memuat permohonan merek Monotes nomor DID2023094257, diterima 16 Oktober 2023 pukul 14.27.43, dengan PT Bintang Mono Indonesia sebagai pemohon. Nama dan alamat muncul pada halaman 1607; lanjutan pada halaman 1608 mencantumkan kelas 10 dan uraian yang mencakup rapid test, alat pemantau diabetes, peralatan diagnostik medis, serta tabung kapiler darah. Ini bukti publikasi permohonan, bukan konfirmasi bahwa merek sudah terdaftar atau tetap aktif. [10]')
table(['Hal yang dinilai','Kesimpulan riset','Batas bukti'],[
('Identitas dan alamat','Konsisten pada kanal perusahaan dan publikasi DJKI','Belum memeriksa akta serta dokumen perusahaan terbaru. [1][3][10]'),
('Cakupan portofolio','Terlihat pada katalog perusahaan','Tidak memastikan seluruh barang tersedia atau masih dipasarkan. [4]–[7]'),
('Hubungan dengan mitra','Dicantumkan pada situs Axion','Kontrak dan status hubungan terkini belum diperiksa. [11]'),
('CDAKB','Katalog yang terindeks memuat klaim sertifikasi','Dokumen sertifikat aktif dan masa berlaku belum diperiksa. [13]'),
('Skala bisnis dan layanan','Belum dapat disimpulkan','Omzet, laba, volume penjualan, waktu layanan dan tingkat komplain tidak tersedia dalam bukti yang diperiksa.')
],[1.45,2.25,3.15])
h('Klaim yang tidak dijadikan kesimpulan')
p('Istilah utama, terkemuka, berkualitas, dan harga terbaik pada materi promosi tidak digunakan sebagai bukti peringkat pasar, mutu seluruh produk, atau keunggulan harga. Testimoni pada situs perusahaan belum menggantikan survei pelanggan independen. Klaim sertifikasi dari katalog historis tidak digunakan untuk menyimpulkan kepatuhan seluruh produk saat ini. [2][13][14]')
p('Riset ini tidak menemukan dasar yang cukup untuk menetapkan pangsa pasar, valuasi, kepemilikan akhir, omzet, atau profitabilitas. Ketiadaan data pada sumber yang diperiksa juga tidak berarti perusahaan tidak memiliki dokumen tersebut.')

page('Analisis bisnis dan prioritas pengembangan')
p('Bagian ini merupakan analisis berdasarkan peta publik. Penilaian belum memakai laporan keuangan, data pelanggan, atau perbandingan harga pesaing.')
h('Kekuatan yang dapat dikembangkan')
p('Katalog menampilkan beberapa kebutuhan yang saling berdekatan dalam kegiatan laboratorium. Secara komersial, hal ini memberi titik awal untuk menawarkan lebih dari satu kategori kepada akun yang sama. Identitas dan kontak yang konsisten juga memudahkan calon pembeli mengenali pemasok. Manfaat tersebut perlu dibuktikan melalui tingkat konversi dan pembelian ulang. [1][4][6][7]')
h('Kesenjangan yang terlihat dari luar')
p('Sejumlah halaman produk yang diperiksa menampilkan nama, gambar, dan tautan katalog, sementara rincian teknis tidak muncul sebagai teks halaman. Calon pembeli berpotensi membutuhkan langkah tambahan untuk memeriksa spesifikasi atau dokumen. Ini adalah observasi pada halaman yang diperiksa, bukan audit menyeluruh atas situs atau lampirannya. [8][12]')
h('Peluang yang paling masuk akal untuk diuji')
p('Pertama, buat lembar produk yang menyatukan kode barang, merek, produsen, kemasan, izin dan masa berlaku, penyimpanan, serta dokumen teknis yang disetujui. Tujuannya mempercepat jawaban atas pertanyaan pengadaan dan mengurangi ketidakkonsistenan informasi.')
p('Kedua, pisahkan penawaran menurut pelanggan laboratorium, fasilitas kesehatan pengguna alat, dan subdistributor. Uji apakah masing-masing segmen membutuhkan kecepatan pasokan, dukungan teknis, kelengkapan dokumen, atau ketentuan komersial yang berbeda. Segmentasi ini adalah usulan riset pelanggan, bukan klasifikasi pelanggan BMI yang sudah terbukti.')
p('Ketiga, untuk barang yang sesuai, uji pengingat pemesanan berdasarkan riwayat konsumsi. Hindari mengasumsikan semua produk memiliki pola pemakaian yang sama. Ukur manfaat dari pesanan berulang dan pengurangan kekosongan stok, bukan hanya jumlah pesan yang dikirim.')
h('Risiko yang perlu masuk pengujian')
p('Hipotesis risiko mencakup ketergantungan pada pemasok impor, masa kedaluwarsa barang, konsentrasi pelanggan, dan kemampuan memenuhi dukungan setelah penjualan. Laporan ini belum menemukan ukuran eksposur BMI terhadap masing-masing risiko. Pengujian memerlukan data pemasok, umur persediaan, penjualan, dan catatan layanan.')
p('Analisis daya saing harga belum dilakukan karena belum tersedia pembandingan produk setara dengan spesifikasi, tanggal, syarat pembayaran, dan biaya pengiriman yang sama. Label harga bersaing belum cukup untuk menetapkan keunggulan komersial.')

page('Rencana tindak lanjut dan metode riset')
h('Usulan kerja selama 30 hari')
table(['Periode','Pekerjaan yang diusulkan','Bukti penyelesaian'],[
('Hari 1–7','Kumpulkan dokumen perusahaan dan daftar produk prioritas; cocokkan peran produsen serta distributor.','Daftar produk dan dokumen pendukung dengan pemilik data serta tanggal pemeriksaan.'),
('Hari 8–14','Wawancarai sampel pelanggan tiap segmen; petakan kendala pemesanan dan dukungan.','Catatan kebutuhan serta pola masalah; tidak menggeneralisasi sampel kecil.'),
('Hari 15–21','Uji dua bentuk penawaran pada kelompok pelanggan sebanding.','Catatan jumlah penawaran, respons, pesanan, margin dan alasan penolakan.'),
('Hari 22–30','Tinjau hasil serta perbarui materi produk dan proses layanan.','Keputusan melanjutkan, mengubah, atau menghentikan setiap uji.')
],[.9,3.05,2.9])
p('Indikator yang disarankan: waktu respons median, proporsi penawaran menjadi pesanan, nilai margin setelah ongkos terkait, pemenuhan pesanan, pembelian ulang, dan kelengkapan dokumen. Tetapkan definisi serta baseline sebelum uji; laporan ini tidak menetapkan target angka tanpa data awal.')
h('Penerapan autoresearch')
p('Alur autoresearch disesuaikan menjadi siklus pertanyaan, pencarian, verifikasi, sintesis, dan pemeriksaan ulang. Putaran pertama memeriksa identitas; putaran kedua memetakan produk dan kanal; putaran ketiga menguji batas klaim melalui publikasi DJKI, situs mitra, dan akses katalog. Tidak ada eksperimen laboratorium atau pelatihan model dalam pekerjaan ini.')
p('H1 tentang identitas bisnis didukung oleh konsistensi beberapa sumber. H2 tentang portofolio didukung sebatas barang yang dipublikasikan perusahaan. H3 tentang peluang pengembangan menghasilkan hipotesis komersial untuk diuji, belum merupakan bukti kenaikan penjualan. Catatan penelitian disimpan agar langkah verifikasi dapat dilanjutkan.')
h('Keterbatasan dan penggunaan')
p('Tanggal acuan laporan adalah 30 Oktober 2026. Halaman web dan profil dapat berubah. Mesin pencari mengembalikan sebagian konten dari indeks; beberapa halaman langsung gagal dibaca. Publikasi DJKI diperiksa dari PDF yang diunduh, termasuk halaman relevan. Dokumen ini ditujukan untuk orientasi bisnis dan penyusunan pertanyaan lanjutan, bukan pengesahan legalitas, audit mutu, atau evaluasi klinis produk.')

page('Daftar sumber')
p('Nomor sumber pada isi mengacu pada daftar berikut. Seluruh sumber diakses atau ditelusuri pada 30 Oktober 2026. Kanal perusahaan diberi label agar berbeda dari bukti pihak lain.')
sources=[
('1','BMI','Hubungi Kami','https://bintangmonoindonesia.com/contact-us/','Kanal perusahaan; alamat dan kontak.'),
('2','BMI','Tentang Kami','https://bintangmonoindonesia.com/about-us/','Kanal perusahaan; profil dan nilai. Konten indeks pencarian; akses langsung mengalami timeout.'),
('3','LinkedIn BMI','Profil perusahaan','https://id.linkedin.com/company/pt-bintang-mono-indonesia','Pelaporan perusahaan; tahun berdiri, ukuran dan merek.'),
('4','BMI','Blood Collection Tube','https://bintangmonoindonesia.com/categories/blood-collection-tube/','Katalog perusahaan; contoh tabung dan perlengkapan sampel.'),
('5','BMI','Local Manufacturing AKD','https://bintangmonoindonesia.com/categories/local-manufacturing-akd/','Pengelompokan katalog perusahaan; bukan verifikasi pabrik atau TKDN.'),
('6','BMI','Fluorescence Immunoassay Analyzer','https://bintangmonoindonesia.com/categories/fluorescence-immunoassay-analyzer/','Katalog perusahaan; MONO X dan Mono Hba1c.'),
('7','BMI','Infectious Disease Orient Gene','https://bintangmonoindonesia.com/categories/infectious-disease-orient-gene/','Katalog perusahaan; rapid test.'),
('8','BMI','Alcohol Swab','https://bintangmonoindonesia.com/product/alcohol-swab/','Halaman produk; tautan e-katalog.'),
('9','BMI','Activated Clotting Time ACT adalah','https://bintangmonoindonesia.com/activated-clotting-time-act-adalah/','Artikel 31 Maret 2026; menyebut Hemonart dan tautan Inaproc. Akses Inaproc gagal dengan respons 402.'),
('10','DJKI','Publikasi merek Monotes','https://www.dgip.go.id/uploads/berita_resmi/file/b3bf2d3e7387defa6a5c12c1eca61782.pdf','Sumber pemerintah; PDF 1858 halaman, hlm 1607–1608. Permohonan DID2023094257, bukan bukti status pendaftaran terkini.'),
('11','PT Axion Jaya Alkesindo','Mitra Kerjasama','https://axionjayaalkesindo.com/','Situs pihak lain; mencantumkan BMI sebagai mitra.'),
('12','BMI','MONO X product','https://bintangmonoindonesia.com/product/mono-x-fluorescence-immunoassay-analyzer-immunofluorescencce-chromatography/','Halaman produk; nama, gambar dan tautan katalog.'),
('13','SMS PT','Katalog Healgen','https://sms-pt.com/Brosur/Healgen/KATALOG%20HEALGEN%20copy.pdf','Katalog yang dihosting pihak lain; klaim CDAKB terlihat pada indeks pencarian. File langsung melampaui batas alat; tidak digunakan untuk mengesahkan status aktif.'),
('14','BMI','Beranda','https://bintangmonoindonesia.com/','Kanal perusahaan; testimoni dan tautan Monoklik menuju halaman login.')]

def link(para,label,url):
    rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
    co=OxmlElement('w:color');co.set(qn('w:val'),'245575');prop.append(co)
    run.append(prop); tx=OxmlElement('w:t');tx.text=label;run.append(tx);el.append(run);para._p.append(el)
for num,owner,title,url,note in sources:
    para=doc.add_paragraph()
    para.paragraph_format.space_after=Pt(7)
    para.add_run(f'[{num}] {owner} — ').bold=True
    link(para,title,url)
    para.add_run('. '+note)
    for r in para.runs:r.font.size=Pt(10)

doc.core_properties.title='Riset Bintang Mono Indonesia'
doc.core_properties.subject='Profil usaha dan peluang pengembangan berbasis sumber publik'
doc.core_properties.author='Codex'
doc.save(OUT)
(ROOT/'data'/'sources.json').write_text(json.dumps([dict(id=n,publisher=o,title=t,url=u,note=note) for n,o,t,u,note in sources],ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'findings.md').write_text('# Temuan riset Bintang Mono Indonesia\n\n'+'\n\n'.join(alltext),encoding='utf-8')
(ROOT/'literature'/'survey.md').write_text('# Daftar sumber\n\n'+'\n\n'.join(f'[{n}] {o} — {t}\n{u}\n{note}' for n,o,t,u,note in sources),encoding='utf-8')
print(OUT)
